import os, re, time, docx, shutil, win32com.client, pythoncom
from datetime import datetime
from PIL import Image
from tqdm.auto import tqdm
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from docx import Document 
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

def get_user_input(): 
    email_dict = {
        "e": "e.desinety@sycomore-am.com",
        "t": "tony.lebon@sycomore-am.com",
        "a": "alex.mory@sycomore-am.com"
    }
    company_name = input("Enter the company name as saved in the SSC folder: ")
    deal_entry_code = input("Enter the NetRoadshow Deal Entry Code: ")

    print("Email options:")
    for key, value in email_dict.items():
        print(f"Enter '{key}' for {value}")
    
    email_short = input("NetRoadshow email address: ").lower()
    email = email_dict.get(email_short, email_short)

    mode = int(input("Enter 1 for 'debug mode' (keep Word doc and screenshots) or 2 for 'user mode' (only pdf file): "))
    return company_name, deal_entry_code, email, mode

def initialize_driver():
    options = webdriver.ChromeOptions()
    options.add_argument("start-maximized")
    options.add_argument(
    "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/87.0.4280.88 Safari/537.36")
    driver = webdriver.Chrome(options=options)
    print("Initializing the WebDriver...")
    return driver

def login_to_netroadshow(driver, deal_entry_code, email):
    print("Opening the Net Roadshow login page...")
    login_url = 'https://www.netroadshow.com/nrs/home/'
    driver.get(login_url)
    time.sleep(1)

    driver.find_element(By.ID, 'homeEntryCodeInput').send_keys(deal_entry_code, Keys.RETURN)
    time.sleep(1)
    driver.find_element(By.ID, 'homeEmailInput').send_keys(email, Keys.RETURN)
    time.sleep(5)
    
    if len(driver.window_handles) > 1:
        driver.maximize_window()
        driver.switch_to.window(driver.window_handles[1])
        time.sleep(6)
        try:
            driver.find_element(By.XPATH, "//button[normalize-space()='Slides-Only']").click()
            time.sleep(5)
            driver.switch_to.window(driver.window_handles[2])
            time.sleep(6)
        except NoSuchElementException:
            pass
        
        try:
            driver.find_element(By.CLASS_NAME, 'disclaimer-btn.btn-agree').click()
            time.sleep(3)
        except (NoSuchElementException, TimeoutException):
            print("No element found.")
        return True
    else:
        print("Failed to login, check credentials or NRS expiration.")
        return False

def login_with_url(driver, deal_entry_code, email):
    print("Opening the Net Roadshow login page...")
    driver.get(deal_entry_code)
    time.sleep(1)
    try:
        driver.find_element(By.ID, 'companyEmailAddressInput').send_keys(email, Keys.RETURN)
        time.sleep(5)
        return True
    except (NoSuchElementException, TimeoutException):
        print("No element found.")
        return False

def handle_welcome_back_modal(driver):
    try:
        driver.find_element(By.CLASS_NAME, 'user-input-restart').click()
        time.sleep(3)
    except NoSuchElementException:
        pass


def verify_video_status(driver):
    try:
        media_player_root = WebDriverWait(driver, 10).until(
            EC.visibility_of_element_located((By.ID, 'media-player-root'))
        )
        if media_player_root.size['width'] > 0 and media_player_root.size['height'] > 0:
            ActionChains(driver).move_to_element(media_player_root).click().perform()
            time.sleep(2)
            video_container = driver.find_element(By.CLASS_NAME, 'media-player-main')
            print("Video is paused." if 'is-paused' in video_container.get_attribute('class') else "Video is not paused.")
        else:
            print("Media player root has no size or location.")
    except (NoSuchElementException, TimeoutException):
        print("No video element found.")

def get_nb_slides(driver):
    slides_label_text = driver.find_element(By.CSS_SELECTOR, '.navbar-slides-container .slides-label span').text.strip()
    return int(re.search(r'All (\d+) Slides', slides_label_text).group(1))

def expand_slide_view(driver):
    print("Zooming on the slide")
    ActionChains(driver).move_to_element(driver.find_element(By.CLASS_NAME, 'slide-image-container')).perform()
    time.sleep(1)
    driver.find_element(By.CLASS_NAME, 'zoom-btn-container.full-screen').click()
    time.sleep(3)

def navigate_to_first_slide(driver):
    print("Navigating to the first slide...")
    try:
        left_arrow = driver.find_element(By.CLASS_NAME, 'arrow-left')
        if 'disabled' not in left_arrow.get_attribute('class'):
            while 'disabled' not in left_arrow.get_attribute('class'):
                left_arrow.click()
                time.sleep(1)
                left_arrow = driver.find_element(By.CLASS_NAME, 'arrow-left')
    except NoSuchElementException:
        print("Left arrow not found.")

def create_output_folders(company_name):
    print("Creating output folders for saving screenshots...")
    current_time = datetime.now().strftime('%Y-%m-%d')
    output_folder = os.path.join("F:\\GESTION\\SSC", company_name, f'{current_time} NRS {company_name}')
    screenshots_folder = os.path.join(output_folder, f'screenshots_{company_name}')
    os.makedirs(output_folder, exist_ok=True)
    os.makedirs(screenshots_folder, exist_ok=True)
    return screenshots_folder, output_folder

def take_screenshots (driver, screenshots_folder, nb_slides, progress_callback=None):
    print("Starting roadshow scraping...")
    image_index = 1
    screenshot_paths = []
    driver.set_window_size(1920,1080)
    
    with tqdm(total=nb_slides, desc='Taking screenshots', unit='slide') as pbar:
        while True:
            try:
                img_element = driver.find_element(By.ID, 'slideImg')
                screenshot_path = os.path.join(screenshots_folder, f'img_{image_index}.png')
                img_element.screenshot(screenshot_path)
                screenshot_paths.append(screenshot_path)
                image_index += 1
                pbar.update(1)

                if progress_callback:
                    progress_callback(image_index - 1)
                
                # Check if the right arrow is disabled
                right_arrow = driver.find_element(By.CLASS_NAME, 'arrow-right')
                if 'disabled' in right_arrow.get_attribute('class'):
                    print(f"Scraping completed. {image_index - 1} screenshots saved in {screenshots_folder}")
                    break
                right_arrow.click()
                time.sleep(2)
                
            except (NoSuchElementException, TimeoutException) as e:
                print(f"Error: {e}")
                break

    return screenshot_paths

def save_screenshots_to_word(screenshot_paths, output_folder, company_name):
    print("Starting screenshot compilation in a Word Docx...")
    current_time = datetime.now().strftime('%Y-%m-%d')
    doc = Document()

    with Image.open(screenshot_paths[0]) as img:
        img_width, img_height = img.size
    aspect_ratio = img_width / img_height
    
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    new_width = section.page_width
    new_height = new_width / aspect_ratio
    section.page_height = int(new_height)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
    
    for screenshot_path in screenshot_paths:
        doc.add_picture(screenshot_path, width=new_width, height=new_height)
        
    doc_path = os.path.join(output_folder, f'{current_time} NRS {company_name}.docx')
    doc.save(doc_path)
    print(f'Word document saved as {doc_path}')
    return doc_path

def save_screenshots_to_pdf(doc_path):
    print("Saving docx as pdf...")
    pdf_path = os.path.splitext(doc_path)[0] + '.pdf'
    try:
        word = win32com.client.gencache.EnsureDispatch('Word.Application', pythoncom.CoInitialize())
        doc = word.Documents.Open(doc_path)
        doc.SaveAs(pdf_path, FileFormat=17) #wdFormatPDF
        doc.Close()
        word.Quit()
        print(f'PDF document saved as {pdf_path}')
    except Exception as e:
        print("An error occurred:", e)
    finally:    
        word = None
    return pdf_path

def cleanup_files(output_folder, pdf_path, mode):
    if mode == 2:
        print("Cleaning NRS folder...")

        parent_folder = os.path.dirname(output_folder)
        new_pdf_path = os.path.join(parent_folder, os.path.basename(pdf_path))
        shutil.move(pdf_path, new_pdf_path)
        print(f'Moved PDF to {new_pdf_path}')
        
        if os.path.exists(output_folder):
            shutil.rmtree(output_folder)
            print(f'Deleted {output_folder}')
